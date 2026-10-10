"""Excel report. Good practice for following foreign keys with __ and select_related."""
from datetime import date

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .models import ReturnItem, SaleItem

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _header(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, title in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(title) + 2)
    ws.freeze_panes = "A2"


@login_required
def sales_report_xlsx(request):
    """/reports/sales.xlsx?from=2026-01-01&to=2026-12-31"""
    start = parse_date(request.GET.get("from", "")) or date.today().replace(day=1)
    end = parse_date(request.GET.get("to", "")) or date.today()

    # Filtering THROUGH the FK: SaleItem -> sale -> deleted_at / created_at
    sale_items = (SaleItem.objects
                  .filter(sale__deleted_at__isnull=True,
                          sale__created_at__date__range=(start, end))
                  .select_related("sale", "sale__customer", "tile")  # one query, JOINs
                  .order_by("sale__created_at"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    _header(ws, ["Sale #", "Date", "Customer", "Tile", "Size", "Pieces", "Boxes",
                 "Loose pcs", "m2", "Price $/m2", "Discount %", "Total $", "Total UZS"])
    total_usd = 0
    for item in sale_items:
        d = item.tile.describe_pieces(item.quantity_pieces)
        ws.append([
            item.sale.number,
            timezone.localtime(item.sale.created_at).strftime("%Y-%m-%d %H:%M"),
            item.sale.customer.name if item.sale.customer else "Walk-in",
            item.tile.code, item.tile.size_label,
            item.quantity_pieces, d["boxes"], d["loose_pieces"], float(d["m2"]),
            float(item.price_usd_per_m2), float(item.discount_pct),
            float(item.line_total_usd),
            float(item.line_total_usd * item.sale.exchange_rate),
        ])
        total_usd += float(item.line_total_usd)
    ws.append([])
    ws.append(["TOTAL", "", "", "", "", "", "", "", "", "", "", round(total_usd, 2)])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=12).font = Font(bold=True)

    # Second sheet: ReturnItem -> sale_return, and -> sale_item -> tile / sale
    returns = (ReturnItem.objects
               .filter(sale_return__created_at__date__range=(start, end))
               .select_related("sale_return", "sale_item__tile", "sale_item__sale"))
    ws2 = wb.create_sheet("Returns")
    _header(ws2, ["Return #", "Date", "Original sale", "Tile", "Pieces", "Condition"])
    for r in returns:
        ws2.append([
            r.sale_return.number,
            timezone.localtime(r.sale_return.created_at).strftime("%Y-%m-%d %H:%M"),
            r.sale_item.sale.number, r.sale_item.tile.code,
            r.quantity_pieces, r.get_condition_display(),
        ])

    response = HttpResponse(content_type=XLSX_TYPE)
    response["Content-Disposition"] = f'attachment; filename="sales_{start}_{end}.xlsx"'
    wb.save(response)
    return response
